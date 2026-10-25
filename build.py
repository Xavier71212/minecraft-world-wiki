#!/usr/bin/env python3
"""Minecraft World Wiki 建站程式。

把 content/（本機「已發布」資料夾的鏡像）裡的 Word 條目轉成百科風靜態網站，輸出到 _site/。
- 子資料夾 = 分類，會自動產生分類頁。
- 條目裡的超連結依目標狀態上色：已發布 → 藍色連結；已有條目但未發布（列在 entries.txt）→ 紅字；其他 → 一般文字。
- 文件中的圖片（旗幟、標誌、路線圖）會自動取出。
- 雙語：繁體中文在網站根目錄，英文在 en/。英文條目來自 content-en/（本機「English」資料夾的鏡像），
  與 content/ 相同的路徑與檔名；每一頁右上角都有維基百科式的語言選單。
"""
import hashlib
import html
import json
import os
import re
import shutil
import subprocess
import sys
import urllib.parse
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parent
CONTENT = ROOT / "content"
CONTENT_EN = ROOT / "content-en"
STATIC = ROOT / "static"
OUT = ROOT / "_site"
CONFIG = json.loads((ROOT / "config.json").read_text(encoding="utf-8"))

E = html.escape

LANGS = ("zh", "en")
LANG_ATTR = {"zh": "zh-Hant", "en": "en"}
LANG_NAME = {"zh": "繁體中文", "en": "English"}
PREFIX = {"zh": "", "en": "en/"}

UI = {
    "zh": {
        "skip": "跳到內容", "search_label": "搜尋條目", "search_ph": "搜尋國家、城市、事件…", "search_btn": "搜尋",
        "tools": "工具", "random": "隨機條目", "recent": "最近更新", "about": "關於本站",
        "nav": "導覽", "home": "首頁", "regions_group": "大陸與海洋", "toc": "目錄", "top": "（頂部）",
        "cat": "分類：", "sep": "｜", "updated": "本條目最後更新於 {d}。", "crumbs": "位置",
        "main_entry": "主條目：", "subcats": "子分類", "in_cat": "此分類中的條目（{n}）",
        "cat_empty": "此分類目前還沒有已發布的條目。",
        "welcome": "歡迎來到 {s}", "count": "目前收錄 <b>{n}</b> 個條目", "map_click": "點擊地圖可放大檢視。",
        "map_alt": "世界地圖與各國勢力範圍", "featured": "特色條目", "dyk": "你知道嗎？", "dyk_lead": "……",
        "regions": "國家與地區", "read_more": "（閱讀全文）", "feat_empty": "網站剛剛開張，第一批條目即將發布。",
        "dyk_empty": "敬請期待。", "recent_empty": "目前沒有條目。",
        "search_title": "搜尋", "search_h": "搜尋結果", "searching": "搜尋中……",
        "random_going": "正在帶你前往隨機條目……",
        "nf_title": "找不到頁面", "nf_h": "找不到這個頁面", "nf_body": "這個條目可能尚未發布，或網址有誤。可以試試上方的搜尋。",
        "unpub": "條目尚未發布", "lang": "語言", "lang_menu": "切換語言", "no_version": "尚無此語言版本",
        "see_also": "參見", "colon": "：",
    },
    "en": {
        "skip": "Skip to content", "search_label": "Search articles", "search_ph": "Search countries, towns, events…", "search_btn": "Search",
        "tools": "Tools", "random": "Random article", "recent": "Recent changes", "about": "About",
        "nav": "Navigation", "home": "Main page", "regions_group": "Continents and oceans", "toc": "Contents", "top": "(Top)",
        "cat": "Categories: ", "sep": " | ", "updated": "This page was last edited on {d}.", "crumbs": "Location",
        "main_entry": "Main article: ", "subcats": "Subcategories", "in_cat": "Pages in this category ({n})",
        "cat_empty": "There are no articles in this category in English yet.",
        "welcome": "Welcome to {s}", "count": "<b>{n}</b> articles in English", "map_click": "Click the map to enlarge it.",
        "map_alt": "World map showing the spheres of influence of each country", "featured": "Featured article", "dyk": "Did you know…", "dyk_lead": "… ",
        "regions": "Countries and regions", "read_more": " (Full article…)", "feat_empty": "The first English articles will be published soon.",
        "dyk_empty": "Coming soon.", "recent_empty": "No articles yet.",
        "search_title": "Search", "search_h": "Search results", "searching": "Searching…",
        "random_going": "Taking you to a random article…",
        "nf_title": "Page not found", "nf_h": "Page not found", "nf_body": "This article may not have been published yet, or the address may be wrong. Try the search box above.",
        "unpub": "This article has not been published yet", "lang": "Language", "lang_menu": "Change language", "no_version": "Not available in this language yet",
        "see_also": "See also", "colon": ": ",
        "ill": "Not yet available in English — opens the Chinese article",
        "stale": "This translation is based on an earlier version of the Chinese article and may be out of date. ",
        "stale_link": "Read the latest Chinese version",
        "zh_name": "Chinese",
    },
}

CATEGORY_CLASS = {
    "國家與地區": "cat-region",
    "交通": "cat-transport",
    "國際組織": "cat-org",
    "宗教": "cat-religion",
    "世界總覽": "cat-history",
    "歷史": "cat-history",
    "典籍": "cat-history",
    "人物": "cat-history",
}

BULLET_RE = re.compile(r"^\s*[•●・‧▪◦]\s*")
SEE_ALSO = ("參見", "See also")

ENT = {}  # 中文條目名稱 -> 資訊（main() 中設定）


def slug(name):
    """檔名安全的網址片段（保留中文，GitHub Pages 可正常提供）。"""
    s = re.sub(r'[\\/:*?"<>|#%]', "-", name).strip()
    return s or "entry"


def slug_l(name, lang):
    """英文網址用底線取代空格（維基百科式，例如 Yerik_Levitch）；中文網址維持原樣。"""
    return slug(name).replace(" ", "_") if lang == "en" else slug(name)


def url_quote(path):
    return urllib.parse.quote(path, safe="/.-_~")


def cfg(lang, key, default=None):
    if lang == "en":
        v = CONFIG.get("en", {}).get(key)
        if v is not None:
            return v
    return CONFIG.get(key, default)


def cat_name(c, lang):
    """分類（資料夾）名稱在該語言的顯示名稱。"""
    if lang == "zh":
        return c
    m = CONFIG.get("en", {}).get("categories", {})
    if c in m:
        return m[c]
    if c in ENT and ENT[c].get("en"):
        return ENT[c]["en"]["title"]
    return c


def sha1(path):
    return hashlib.sha1(Path(path).read_bytes()).hexdigest()


# ---------------------------------------------------------------- 收集條目

def collect():
    entries = {}  # stem -> info
    for f in sorted(CONTENT.rglob("*.docx")):
        if f.name.startswith("~$"):
            continue
        rel = f.relative_to(CONTENT)
        stem = f.stem.strip()
        if stem in entries:
            print(f"警告：條目名稱重複，略過 {rel}", file=sys.stderr)
            continue
        entries[stem] = {
            "stem": stem,
            "file": f,
            "rel": rel.as_posix(),
            "cats": list(rel.parent.parts),
            "url": f"wiki/{url_quote(slug(stem))}.html",
            "out": f"wiki/{slug(stem)}.html",
        }
    # 英文版：content-en/ 中與中文條目相同路徑的檔案
    used = set()
    for stem, info in entries.items():
        ef = CONTENT_EN / info["rel"]
        if not ef.exists():
            continue
        try:
            d = Document(str(ef))
        except Exception as e:
            print(f"無法讀取英文條目 {ef.name}：{e}", file=sys.stderr)
            continue
        title = (d.core_properties.title or "").strip()
        if not title:
            for p in d.paragraphs:
                if p.style is not None and p.style.name.lower() in ("title", "標題") and p.text.strip():
                    title = p.text.strip()
                    break
        title = title or stem
        if title in used:
            title = f"{title} ({stem})"
        used.add(title)
        m = re.search(r"source-sha1:([0-9a-f]{40})", d.core_properties.comments or "")
        info["en"] = {
            "file": ef,
            "title": title,
            "rel": f"wiki/{url_quote(slug_l(title, 'en'))}.html",
            "url": f"en/wiki/{url_quote(slug_l(title, 'en'))}.html",
            "out": f"en/wiki/{slug_l(title, 'en')}.html",
            "stale": bool(m) and m.group(1) != sha1(info["file"]),
        }
    existing = set()
    ex_file = ROOT / "entries.txt"
    if ex_file.exists():
        existing = {l.strip() for l in ex_file.read_text(encoding="utf-8").splitlines() if l.strip()}
    return entries, existing


def git_dates():
    """每個檔案最後一次提交的時間（最近更新用）。"""
    dates = {}
    try:
        out = subprocess.run(
            ["git", "log", "--name-only", "--format=@@%cI", "--", "content", "content-en"],
            cwd=ROOT, capture_output=True, text=True, check=True,
            env={**os.environ, "GIT_CONFIG_PARAMETERS": "'core.quotepath=false'"},
        ).stdout
        cur = None
        for line in out.splitlines():
            if line.startswith("@@"):
                cur = line[2:]
            elif line.strip() and cur:
                dates.setdefault(line.strip(), cur)
    except Exception as e:  # 本機沒有 git 時不影響建站
        print("無法讀取 git 紀錄：", e, file=sys.stderr)
    return dates


def entry_url(stem, lang):
    """條目在該語言的網址（相對網站根目錄）；英文版沒有翻譯時回傳 None。"""
    info = ENT[stem]
    if lang == "en":
        return info["en"]["url"] if info.get("en") else None
    return info["url"]


def entry_title(stem, lang):
    info = ENT[stem]
    return info["en"]["title"] if lang == "en" and info.get("en") else stem


# ---------------------------------------------------------------- docx 轉 HTML

class Converter:
    def __init__(self, entries, existing, media_dir, lang="zh"):
        self.entries = entries
        self.existing = existing
        self.media_dir = media_dir
        self.lang = lang
        self.media_dir.mkdir(parents=True, exist_ok=True)

    def link_state(self, target, text):
        stem = None
        if target:
            t = urllib.parse.unquote(target).replace("\\", "/").split("#")[0]
            name = t.rsplit("/", 1)[-1]
            if name.lower().endswith(".docx"):
                stem = name[:-5].strip()
        for cand in (stem, text.strip()):
            if cand and cand in self.entries:
                return "pub", cand
        for cand in (stem, text.strip()):
            if cand and cand in self.existing:
                return "unpub", cand
        return "none", None

    def image(self, part, rid):
        try:
            img = part.related_parts[rid]
        except KeyError:
            return None
        blob = img.blob
        ext = Path(img.partname).suffix.lower() or ".png"
        if ext in (".emf", ".wmf"):
            return None
        name = hashlib.sha1(blob).hexdigest()[:16] + ext
        dest = self.media_dir / name
        if not dest.exists():
            dest.write_bytes(blob)
        return name

    def link_html(self, stem, inner, base):
        u = UI[self.lang]
        url = entry_url(stem, self.lang)
        if url:
            return f'<a href="{base}{url}">{inner}</a>'
        # 英文頁面：目標條目還沒有英文版 → 連到中文條目並標示 zh
        return (f'<a class="ill" href="{base}{self.entries[stem]["url"]}" hreflang="zh-Hant" title="{E(u["ill"])}">{inner}</a>'
                f'<sup class="ill-tag" lang="en">zh</sup>')

    def runs_html(self, el, part, base, imgs):
        out = []
        for c in el.iterchildren():
            tag = c.tag.split("}")[1]
            if tag == "r":
                out.append(self.run_html(c, part, base, imgs))
            elif tag == "hyperlink":
                text = "".join(x.text or "" for x in c.iter(qn("w:t")))
                rid = c.get(qn("r:id"))
                target = None
                if rid and rid in part.rels:
                    target = part.rels[rid].target_ref
                inner = "".join(self.run_html(r, part, base, imgs) for r in c.iter(qn("w:r")))
                state, stem = self.link_state(target, text)
                if state == "pub":
                    out.append(self.link_html(stem, inner, base))
                elif state == "unpub":
                    out.append(f'<span class="unpub" title="{E(UI[self.lang]["unpub"])}">{inner}</span>')
                elif target and re.match(r"https?://", target):
                    out.append(f'<a href="{E(target)}" rel="noopener">{inner}</a>')
                else:
                    out.append(inner)
            elif tag in ("smartTag", "ins", "fldSimple"):
                out.append(self.runs_html(c, part, base, imgs))
        return "".join(out)

    def run_html(self, r, part, base, imgs):
        parts = []
        rpr = r.find(qn("w:rPr"))
        bold = rpr is not None and rpr.find(qn("w:b")) is not None and rpr.find(qn("w:b")).get(qn("w:val")) not in ("0", "false")
        ital = rpr is not None and rpr.find(qn("w:i")) is not None and rpr.find(qn("w:i")).get(qn("w:val")) not in ("0", "false")
        for c in r.iterchildren():
            tag = c.tag.split("}")[1]
            if tag == "t":
                parts.append(E(c.text or ""))
            elif tag == "tab":
                parts.append("　")
            elif tag == "br":
                parts.append("<br>")
            elif tag == "drawing":
                for blip in c.iter(qn("a:blip")):
                    name = self.image(part, blip.get(qn("r:embed")))
                    if name:
                        imgs.append(name)
        s = "".join(parts)
        if s and bold:
            s = f"<b>{s}</b>"
        if s and ital:
            s = f"<i>{s}</i>"
        return s

    def convert(self, path, base):
        doc = Document(str(path))
        part = doc.part
        body = doc.element.body
        blocks = []  # (kind, data)
        infobox = None
        seen_heading = False
        for el in body.iterchildren():
            tag = el.tag.split("}")[1]
            if tag == "p":
                p = Paragraph(el, doc)
                style = (p.style.name if p.style is not None else "") or ""
                imgs = []
                h = self.runs_html(el, part, base, imgs).strip()
                text = p.text.strip()
                if style in ("Title", "標題") or style.lower() == "title":
                    continue
                if style.startswith("Heading") or style.startswith("標題"):
                    lvl = 3 if ("3" in style or "4" in style) else 2
                    if text:
                        blocks.append(("h", (lvl, text)))
                        seen_heading = True
                    continue
                for name in imgs:
                    blocks.append(("img", name))
                if not text:
                    continue
                if BULLET_RE.match(text):
                    h = BULLET_RE.sub("", re.sub(r"^(\s|&nbsp;)*", "", h), count=1)
                    h = re.sub(r"^(<[^>]+>)*\s*[•●・‧▪◦]\s*", lambda m: (m.group(1) or ""), h)
                    blocks.append(("li", h.strip()))
                else:
                    blocks.append(("p", h))
            elif tag == "tbl":
                rows = []
                for tr in el.iter(qn("w:tr")):
                    cells = []
                    for tc in tr.findall(qn("w:tc")):
                        imgs = []
                        parts = []
                        for pe in tc.iter(qn("w:p")):
                            ph = self.runs_html(pe, part, base, imgs).strip()
                            if ph:
                                parts.append(ph)
                        cells.append({"html": "<br>".join(parts), "imgs": imgs,
                                      "text": "".join(x.text or "" for x in tc.iter(qn("w:t"))).strip()})
                    rows.append(cells)
                if infobox is None and not seen_heading and not any(k in ("p", "li") for k, _ in blocks):
                    infobox = rows
                else:
                    blocks.append(("table", rows))
        return blocks, infobox


# ---------------------------------------------------------------- 版面

def cat_class(cats):
    return CATEGORY_CLASS.get(cats[0], "cat-history") if cats else "cat-history"


def lang_menu(lang, alt, base):
    """維基百科式語言選單：目前語言打勾，其他語言連到對應頁面；沒有該語言版本時顯示為灰色。"""
    u = UI[lang]
    items = []
    for l in LANGS:
        name = E(LANG_NAME[l])
        la = LANG_ATTR[l]
        if l == lang:
            items.append(f'<li><span class="cur" lang="{la}" aria-current="true">{name}</span></li>')
        elif alt.get(l):
            items.append(f'<li><a href="{base}{alt[l]}" lang="{la}" hreflang="{la}">{name}</a></li>')
        else:
            items.append(f'<li><span class="na" lang="{la}">{name}<small lang="{LANG_ATTR[lang]}">{E(u["no_version"])}</small></span></li>')
    return (
        f'<details class="langsw"><summary title="{E(u["lang_menu"])}">'
        f'<svg class="langicon" viewBox="0 0 20 20" width="18" height="18" aria-hidden="true"><path fill="currentColor" d="M20 18h-1.44a.6.6 0 0 1-.4-.12.8.8 0 0 1-.23-.31L17 15h-5l-1 2.54a.8.8 0 0 1-.22.3.6.6 0 0 1-.4.14H9l4.55-11.47h1.89zm-3.53-4.31L14.89 9.5a12 12 0 0 1-.39-1.24q-.09.37-.19.69l-.19.56-1.58 4.19zm-6.3-1.58a13.4 13.4 0 0 1-2.91-1.41 11.46 11.46 0 0 0 2.81-5.37H12V4H7.31a4 4 0 0 0-.2-.56C6.87 2.79 6.6 2 6.6 2l-1.47.5s.4.89.6 1.5H0v1.33h2.15A11.23 11.23 0 0 0 5 10.7a17.2 17.2 0 0 1-5 2.1q.56.82.87 1.38a23.3 23.3 0 0 0 5.22-2.51 15.6 15.6 0 0 0 3.56 1.77zM3.63 5.33h4.91a8.1 8.1 0 0 1-2.45 4.45 9.1 9.1 0 0 1-2.46-4.45"/></svg>'
        f'<span class="langlabel">{E(u["lang"])}</span><span class="caret" aria-hidden="true"></span></summary>'
        f'<ul class="langlist">{"".join(items)}</ul></details>'
    )


def page(title, body, base, *, lang="zh", alt=None, sidebar="", description="", page_class=""):
    u = UI[lang]
    alt = alt or {}
    site = cfg(lang, "site_title")
    full = site if title == site else f"{title} - {site}"
    lr = base + PREFIX[lang]  # 該語言版本的根目錄
    site_url = CONFIG.get("site_url", "").rstrip("/") + "/"
    alternates = "".join(
        f'<link rel="alternate" hreflang="{LANG_ATTR[l]}" href="{E(site_url + alt[l])}">' for l in LANGS if alt.get(l)
    )
    return f"""<!doctype html>
<html lang="{LANG_ATTR[lang]}">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{E(full)}</title>
<meta name="description" content="{E(description or cfg(lang, 'tagline'))}">
{alternates}
<link rel="preconnect" href="https://fonts.googleapis.com">
<link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+TC:wght@400;500;700&family=Noto+Serif+TC:wght@500;700&display=swap" rel="stylesheet">
<link rel="stylesheet" href="{base}static/style.css">
<link rel="icon" href="{base}static/favicon.svg" type="image/svg+xml">
</head>
<body class="{page_class}" data-base="{lr}">
<a class="skip" href="#main">{E(u['skip'])}</a>
<header class="topbar">
  <div class="wrap topbar-in">
    <a class="brand" href="{lr}index.html">
      <span class="mark" aria-hidden="true">世</span>
      <span class="brand-text"><span class="brand-name">{E(site)}</span><span class="brand-sub">{E(cfg(lang, 'short_tagline'))}</span></span>
    </a>
    <form class="search" role="search" action="{lr}search.html" method="get">
      <label for="q" class="sr">{E(u['search_label'])}</label>
      <input id="q" name="q" type="search" placeholder="{E(u['search_ph'])}" autocomplete="off">
      <button type="submit">{E(u['search_btn'])}</button>
      <ul class="suggest" hidden></ul>
    </form>
    <nav class="toplinks" aria-label="{E(u['tools'])}">
      <a href="{lr}random.html">{E(u['random'])}</a>
      <a href="{lr}recent.html">{E(u['recent'])}</a>
      <a href="{lr}about.html">{E(u['about'])}</a>
    </nav>
  </div>
</header>
<div class="wrap layout">
  <aside class="side">{sidebar}</aside>
  <main id="main" class="main"><div class="pagetools">{lang_menu(lang, alt, base)}</div>{body}</main>
</div>
<footer class="foot"><div class="wrap">{E(cfg(lang, 'footer'))}</div></footer>
<script src="{base}static/site.js"></script>
</body>
</html>
"""


def cat_url(cats, lang="zh"):
    return PREFIX[lang] + "category/" + "/".join(url_quote(slug_l(cat_name(c, lang), lang)) for c in cats) + "/index.html"


def cat_out(cats, lang="zh"):
    return Path(PREFIX[lang] or ".") / "category" / Path(*[slug_l(cat_name(c, lang), lang) for c in cats]) / "index.html"


def nav_sidebar(tree, base, lang="zh"):
    u = UI[lang]
    items = "".join(
        f'<a href="{base}{cat_url([c], lang)}">{E(cat_name(c, lang))}</a>' for c in tree.get("dirs", {})
    )
    regions = tree.get("dirs", {}).get("國家與地區", {}).get("dirs", {})
    reg = "".join(
        f'<a href="{base}{cat_url(["國家與地區", r], lang)}">{E(cat_name(r, lang))}</a>'
        for r in regions
    )
    html_ = f'<div class="side-group"><div class="side-label">{E(u["nav"])}</div><a href="{base}{PREFIX[lang]}index.html">{E(u["home"])}</a>{items}</div>'
    if reg:
        html_ += f'<div class="side-group"><div class="side-label">{E(u["regions_group"])}</div>{reg}</div>'
    return html_


def build_tree(entries):
    tree = {"dirs": {}, "entries": []}
    for d in sorted(p for p in CONTENT.rglob("*") if p.is_dir()):
        node = tree
        for part in d.relative_to(CONTENT).parts:
            node = node["dirs"].setdefault(part, {"dirs": {}, "entries": []})
    for info in entries.values():
        node = tree
        for part in info["cats"]:
            node = node["dirs"].setdefault(part, {"dirs": {}, "entries": []})
        node["entries"].append(info["stem"])
    return tree


def visible(stem, lang):
    return lang == "zh" or bool(ENT[stem].get("en"))


def count_entries(node, lang="zh"):
    return sum(1 for e in node["entries"] if visible(e, lang)) + sum(count_entries(n, lang) for n in node["dirs"].values())


def render_blocks(blocks, base):
    out = []
    i = 0
    toc = []
    hid = 0
    while i < len(blocks):
        kind, data = blocks[i]
        if kind == "h":
            lvl, text = data
            hid += 1
            anchor = f"s{hid}"
            toc.append((lvl, text, anchor))
            out.append(f'<h{lvl} id="{anchor}">{E(text)}</h{lvl}>')
        elif kind == "p":
            out.append(f"<p>{data}</p>")
        elif kind == "li":
            items = []
            while i < len(blocks) and blocks[i][0] == "li":
                items.append(f"<li>{blocks[i][1]}</li>")
                i += 1
            cls = ' class="cols"' if toc and toc[-1][1] in SEE_ALSO else ""
            out.append(f"<ul{cls}>{''.join(items)}</ul>")
            continue
        elif kind == "img":
            # 圖片後面緊接的短段落視為圖說
            cap = ""
            if i + 1 < len(blocks) and blocks[i + 1][0] == "p" and len(re.sub("<[^>]+>", "", blocks[i + 1][1])) <= 60:
                cap = blocks[i + 1][1]
                i += 1
            out.append(
                f'<figure class="fig"><a href="{base}media/{data}"><img src="{base}media/{data}" alt="{E(re.sub("<[^>]+>", "", cap))}" loading="lazy"></a>'
                + (f"<figcaption>{cap}</figcaption>" if cap else "") + "</figure>"
            )
        elif kind == "table":
            rows = []
            for r in data:
                cells = ""
                for c in r:
                    pics = "".join('<img src="' + base + "media/" + n + '" alt="">' for n in c["imgs"])
                    cells += "<td>" + c["html"] + pics + "</td>"
                rows.append(f"<tr>{cells}</tr>")
            out.append(f'<div class="tablewrap"><table class="wikitable"><tbody>{"".join(rows)}</tbody></table></div>')
        i += 1
    return "\n".join(out), toc


def render_infobox(rows, title, cats, base, lang="zh", subtitle=""):
    if not rows:
        return ""
    cap = rows[0][0]["text"] if rows[0] and len(rows[0]) == 1 else title
    start = 1 if rows[0] and len(rows[0]) == 1 else 0
    sub_text = subtitle
    trs = []
    for r in rows[start:]:
        imgs = [n for c in r for n in c["imgs"]]
        if imgs:
            labels = " ".join(c["text"] for c in r).split()
            figs = []
            for k, n in enumerate(imgs):
                lab = labels[k] if k < len(labels) else ""
                figs.append(
                    f'<figure><img src="{base}media/{n}" alt="{E(cap + " " + lab if lang == "en" else cap + lab)}"><figcaption>{E(lab)}</figcaption></figure>'
                )
            trs.append(f'<tr><td colspan="2" class="ib-img"><div class="ib-figs">{"".join(figs)}</div></td></tr>')
        elif len(r) >= 2:
            k, v = r[0], r[1]
            if lang == "zh" and k["text"] in ("英文名稱", "英文站名") and not sub_text:
                sub_text = v["text"]
            trs.append(f'<tr><th scope="row">{k["html"]}</th><td>{v["html"]}</td></tr>')
        elif len(r) == 1 and r[0]["text"]:
            trs.append(f'<tr><th colspan="2" class="ib-sub">{r[0]["html"]}</th></tr>')
    sub_lang = ' lang="zh-Hant"' if lang == "en" else ""
    sub = f"<span{sub_lang}>{E(sub_text)}</span>" if sub_text and sub_text != cap else ""
    return (
        f'<table class="infobox"><caption class="{cat_class(cats)}">{E(cap)}{sub}</caption>'
        f'<tbody>{"".join(trs)}</tbody></table>'
    )


def first_text(blocks, limit=120, lang="zh"):
    for k, d in blocks:
        if k == "p":
            t = re.sub("<[^>]+>", "", d)
            t = html.unescape(t).strip()
            if t:
                if lang == "en":
                    limit *= 2
                    return t if len(t) <= limit else t[:limit].rsplit(" ", 1)[0].rstrip(",;:") + "…"
                return t if len(t) <= limit else t[:limit].rstrip() + "……"
    return ""


def breadcrumb(cats, base, current=None, lang="zh"):
    parts = []
    for i in range(len(cats)):
        parts.append(f'<a href="{base}{cat_url(cats[:i + 1], lang)}">{E(cat_name(cats[i], lang))}</a>')
    if current:
        parts.append(E(current))
    return f'<nav class="crumbs" aria-label="{E(UI[lang]["crumbs"])}">' + "　›　".join(parts) + "</nav>" if parts else ""


def write(rel, text):
    out = OUT / rel
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")


def both(path):
    """非條目頁面（首頁、分類、搜尋等）在兩種語言都存在。"""
    return {"zh": path, "en": "en/" + path}


# ---------------------------------------------------------------- 單一語言版本

def build_lang(lang, entries, existing, tree, dates):
    u = UI[lang]
    conv = Converter(entries, existing, OUT / "media", lang)
    depth0 = 1 if lang == "en" else 0  # en/ 多一層
    index = []
    root = "../" * depth0  # 該語言首頁相對網站根目录

    # 條目頁
    for stem, info in entries.items():
        if lang == "en" and not info.get("en"):
            continue
        base = "../" * (depth0 + 1)
        src = info["en"]["file"] if lang == "en" else info["file"]
        title = entry_title(stem, lang)
        try:
            blocks, infobox = conv.convert(src, base)
        except Exception as e:
            print(f"無法轉換 {src.name}：{e}", file=sys.stderr)
            continue
        loc = info["en"] if lang == "en" else info
        loc["summary"] = first_text(blocks, lang=lang)
        rel = src.relative_to(ROOT).as_posix()
        loc["updated"] = dates.get(rel, "")
        body_html, toc = render_blocks(blocks, base)
        ib = render_infobox(infobox, title, info["cats"], base, lang, subtitle=stem if lang == "en" else "")
        loc["thumb"] = None
        if infobox:
            for r in infobox:
                for c in r:
                    if c["imgs"] and not loc["thumb"]:
                        loc["thumb"] = c["imgs"][0]
        toc_html = ""
        if len(toc) >= 2:
            toc_html = f'<div class="side-group"><div class="side-label">{E(u["toc"])}</div><a href="#top">{E(u["top"])}</a>' + "".join(
                f'<a class="toc-l{l}" href="#{a}">{E(t)}</a>' for l, t, a in toc
            ) + "</div>"
        cats_bar = ""
        if info["cats"]:
            cats_bar = f'<div class="catbar">{E(u["cat"])}' + E(u["sep"]).join(
                f'<a href="{base}{cat_url(info["cats"][:i + 1], lang)}">{E(cat_name(info["cats"][i], lang))}</a>' for i in range(len(info["cats"]))
            ) + "</div>"
        upd = ""
        if loc["updated"]:
            upd = f'<p class="updated">{E(u["updated"].format(d=loc["updated"][:10]))}</p>'
        notice = ""
        if lang == "en" and info["en"].get("stale"):
            notice = f'<p class="hatnote notice">{E(u["stale"])}<a href="{base}{info["url"]}" hreflang="zh-Hant">{E(u["stale_link"])}</a>.</p>'
        crumb_cats = info["cats"][:-1] if info["cats"] and info["cats"][-1] == stem else info["cats"]
        body = (
            f'<article class="article" id="top">{breadcrumb(crumb_cats, base, title, lang)}'
            f'<h1 class="title">{E(title)}</h1>{notice}'
            f'<div class="article-body">{ib}{body_html}</div>{cats_bar}{upd}</article>'
        )
        alt = {"zh": info["url"], "en": info["en"]["url"] if info.get("en") else None}
        write(loc["out"], page(title, body, base, lang=lang, alt=alt,
                               sidebar=toc_html or nav_sidebar(tree, base, lang), description=loc["summary"]))
        idx_url = loc["rel"] if lang == "en" else info["url"]
        index.append({"t": title, "a": stem if lang == "en" else (info["en"]["title"] if info.get("en") else ""),
                      "u": idx_url, "s": loc["summary"], "c": " › ".join(cat_name(c, lang) for c in info["cats"])})

    def loc_of(stem):
        return entries[stem]["en"] if lang == "en" else entries[stem]

    # 分類頁
    def walk(node, cats):
        if cats:
            base = "../" * (len(cats) + 1 + depth0)
            main_entry = cats[-1] if cats[-1] in entries and visible(cats[-1], lang) else None
            name = cat_name(cats[-1], lang)
            parts = [f'<article class="article">{breadcrumb(cats[:-1], base, name, lang)}<h1 class="title">{E(name)}</h1>']
            if main_entry:
                parts.append(
                    f'<p class="hatnote">{E(u["main_entry"])}<a href="{base}{entry_url(main_entry, lang)}">{E(entry_title(main_entry, lang))}</a></p>'
                )
            if node["dirs"]:
                parts.append(f"<h2>{E(u['subcats'])}</h2><ul class=\"cols\">" + "".join(
                    f'<li><a href="{base}{cat_url(cats + [d], lang)}">{E(cat_name(d, lang))}</a>（{count_entries(n, lang)}）</li>'
                    if lang == "zh" else
                    f'<li><a href="{base}{cat_url(cats + [d], lang)}">{E(cat_name(d, lang))}</a> ({count_entries(n, lang)})</li>'
                    for d, n in node["dirs"].items()
                ) + "</ul>")
            listed = sorted((e for e in node["entries"] if e != main_entry and visible(e, lang)), key=lambda e: entry_title(e, lang))
            if listed:
                parts.append(f"<h2>{E(u['in_cat'].format(n=len(listed)))}</h2><ul class=\"cols\">" + "".join(
                    f'<li><a href="{base}{entry_url(e, lang)}">{E(entry_title(e, lang))}</a></li>' for e in listed
                ) + "</ul>")
            if not node["dirs"] and not listed and not main_entry:
                parts.append(f'<p class="empty">{E(u["cat_empty"])}</p>')
            parts.append("</article>")
            write(cat_out(cats, lang), page(name, "".join(parts), base, lang=lang,
                                            alt={"zh": cat_url(cats, "zh"), "en": cat_url(cats, "en")},
                                            sidebar=nav_sidebar(tree, base, lang)))
        for d, n in node["dirs"].items():
            walk(n, cats + [d])

    walk(tree, [])

    # 首頁
    base = root
    lr = PREFIX[lang]
    shown = [s for s in entries if visible(s, lang) and loc_of(s).get("summary") is not None]
    featured = None
    for t in cfg(lang, "featured", []):
        if t in shown:
            featured = t
            break
    dated = sorted(shown, key=lambda s: loc_of(s).get("updated", ""), reverse=True)
    if not featured and dated:
        featured = dated[0]
    feat_html = f'<p class="empty">{E(u["feat_empty"])}</p>'
    if featured:
        loc = loc_of(featured)
        title = entry_title(featured, lang)
        href = base + entry_url(featured, lang)
        thumb = f'<img class="feat-img" src="{base}media/{loc["thumb"]}" alt="">' if loc.get("thumb") else ""
        summ = loc["summary"]
        rest = summ[len(title):] if summ.startswith(title) else (u["colon"] + summ)
        feat_html = (
            f'<div class="feat">{thumb}<p><a href="{href}"><b>{E(title)}</b></a>'
            f'{E(rest)}<a href="{href}">{E(u["read_more"])}</a></p></div>'
        )
    dyk_items = []
    for item in cfg(lang, "did_you_know", []):
        def rep(m):
            name, _, text = m.group(1).partition("|")
            text = text or (entry_title(name, lang) if name in entries else name)
            if name in entries:
                url = entry_url(name, lang)
                if url:
                    return f'<a href="{base}{url}">{E(text)}</a>'
                return (f'<a class="ill" href="{base}{entries[name]["url"]}" hreflang="zh-Hant" title="{E(u.get("ill", ""))}">{E(text)}</a>'
                        f'<sup class="ill-tag">zh</sup>')
            return E(text)
        parts_ = re.split(r"(\[\[.+?\]\])", item)
        h = "".join(rep(re.match(r"\[\[(.+?)\]\]", p)) if p.startswith("[[") else E(p) for p in parts_)
        dyk_items.append(f"<li>{E(u['dyk_lead'])}{h}</li>")
    dyk_html = f'<ul class="dyk">{"".join(dyk_items)}</ul>' if dyk_items else f'<p class="empty">{E(u["dyk_empty"])}</p>'

    regions = tree["dirs"].get("國家與地區", {"dirs": {}})
    reg_rows = []
    sep = "・" if lang == "zh" else " · "
    for r, n in regions["dirs"].items():
        links = []
        for c, cn in n["dirs"].items():
            if c in entries and visible(c, lang):
                links.append(f'<a href="{base}{entry_url(c, lang)}">{E(entry_title(c, lang))}</a>')
            elif count_entries(cn, lang):
                links.append(f'<a href="{base}{cat_url(["國家與地區", r, c], lang)}">{E(cat_name(c, lang))}</a>')
            else:
                links.append(f'<span class="plain">{E(cat_name(c, lang))}</span>')
        for e in sorted(n["entries"]):
            if e not in n["dirs"] and visible(e, lang):
                links.append(f'<a href="{base}{entry_url(e, lang)}">{E(entry_title(e, lang))}</a>')
        rh = f'<a href="{base}{cat_url(["國家與地區", r], lang)}">{E(cat_name(r, lang))}</a>'
        reg_rows.append(f'<tr><th scope="row">{rh}</th><td>{sep.join(links) or "—"}</td></tr>')
    reg_html = (
        f'<section class="box"><h2 class="box-h">{E(u["regions"])}</h2><div class="tablewrap"><table class="regions"><tbody>{"".join(reg_rows)}</tbody></table></div></section>'
        if reg_rows else ""
    )

    def recent_li(s):
        return f'<li><a href="{base}{entry_url(s, lang)}">{E(entry_title(s, lang))}</a><span>{loc_of(s).get("updated", "")[:10]}</span></li>'

    recent_html = ""
    if dated:
        recent_html = f'<section class="box"><h2 class="box-h">{E(u["recent"])}</h2><ul class="recent">' + "".join(
            recent_li(s) for s in dated[:8]) + "</ul></section>"
    total = len(index)
    site = cfg(lang, "site_title")
    home = f"""
<section class="welcome">
  <h1>{E(u['welcome'].format(s=site))}</h1>
  <p>{E(cfg(lang, 'tagline'))}</p>
  <p class="count">{u['count'].format(n=total)}</p>
</section>
<figure class="box mapfig">
  <a href="{base}static/map-large.jpg"><img src="{base}static/map.jpg" alt="{E(u['map_alt'])}"></a>
  <figcaption>{E(cfg(lang, 'map_caption'))}{E(u['map_click'])}</figcaption>
</figure>
<div class="grid2">
  <section class="box"><h2 class="box-h head-featured">{E(u['featured'])}</h2><div class="box-b">{feat_html}</div></section>
  <section class="box"><h2 class="box-h head-dyk">{E(u['dyk'])}</h2><div class="box-b">{dyk_html}</div></section>
</div>
{reg_html}
{recent_html}
"""
    write(lr + "index.html", page(site, home, base, lang=lang, alt=both("index.html"),
                                  sidebar=nav_sidebar(tree, base, lang), page_class="home"))

    # 最近更新
    rec_all = "".join(recent_li(s) for s in dated[:100]) or f'<li class="empty">{E(u["recent_empty"])}</li>'
    write(lr + "recent.html", page(u["recent"], f'<article class="article"><h1 class="title">{E(u["recent"])}</h1><ul class="recent">{rec_all}</ul></article>',
                                   base, lang=lang, alt=both("recent.html"), sidebar=nav_sidebar(tree, base, lang)))

    # 搜尋、隨機、關於
    write(lr + "search-index.json", json.dumps(index, ensure_ascii=False))
    write(lr + "search.html", page(u["search_title"], f'<article class="article"><h1 class="title">{E(u["search_h"])}</h1><p id="search-status" class="muted">{E(u["searching"])}</p><ul id="search-results" class="results"></ul></article>',
                                   base, lang=lang, alt=both("search.html"), sidebar=nav_sidebar(tree, base, lang), page_class="search-page"))
    write(lr + "random.html", page(u["random"], f'<article class="article"><h1 class="title">{E(u["random"])}</h1><p id="random-status" class="muted">{E(u["random_going"])}</p></article>',
                                   base, lang=lang, alt=both("random.html"), sidebar=nav_sidebar(tree, base, lang), page_class="random-page"))
    about = "".join(f"<p>{E(p)}</p>" for p in cfg(lang, "about", []))
    write(lr + "about.html", page(u["about"], f'<article class="article"><h1 class="title">{E(u["about"])}</h1>{about}</article>',
                                  base, lang=lang, alt=both("about.html"), sidebar=nav_sidebar(tree, base, lang)))
    return total


# ---------------------------------------------------------------- 主程式

def main():
    global ENT
    if OUT.exists():
        shutil.rmtree(OUT)
    OUT.mkdir()
    shutil.copytree(STATIC, OUT / "static")
    (OUT / ".nojekyll").write_text("")

    entries, existing = collect()
    ENT = entries
    tree = build_tree(entries)
    dates = git_dates()
    totals = {lang: build_lang(lang, entries, existing, tree, dates) for lang in LANGS}

    # 404（GitHub Pages 所有語言共用一頁，兩種語言並列）
    root = "/" + CONFIG.get("base_path", "").strip("/") + "/" if CONFIG.get("base_path") else "/"
    zh, en = UI["zh"], UI["en"]
    nf = (f'<article class="article"><h1 class="title">{E(zh["nf_h"])}</h1><p>{E(zh["nf_body"])}</p>'
          f'<h2 lang="en">{E(en["nf_h"])}</h2><p lang="en">{E(en["nf_body"])} <a href="{root}en/index.html">English main page</a></p></article>')
    write("404.html", page(zh["nf_title"], nf, root, lang="zh", alt=both("index.html"), sidebar=""))

    print(f"完成：中文 {totals['zh']} 個條目、英文 {totals['en']} 個條目，輸出到 {OUT}")


if __name__ == "__main__":
    main()
